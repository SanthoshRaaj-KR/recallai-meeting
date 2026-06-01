"""Generate hr_policies.docx for the sandbox."""
from pathlib import Path
import docx

OUT = Path(__file__).parent / "docs" / "hr_policies.docx"
doc = docx.Document()

doc.add_heading("Paid Time Off", level=1)
doc.add_paragraph(
    "Full-time employees accrue 15 days of paid time off per year. PTO must be "
    "requested at least two weeks in advance and approved by the direct manager. "
    "Unused PTO may be carried over up to a maximum of 5 days into the next year."
)

doc.add_heading("Remote Work Policy", level=1)
doc.add_paragraph(
    "Employees are expected to work from the office at least three days per week. "
    "Remote work days must be coordinated with the team to ensure adequate in-office "
    "coverage. Fully remote arrangements require VP approval."
)

doc.add_heading("Expense Reimbursement", level=1)
doc.add_paragraph(
    "Business expenses under 100 dollars may be approved by a direct manager. "
    "Expenses above that threshold require finance approval. All expense reports "
    "must be submitted within 30 days of the expense being incurred."
)

doc.add_heading("Performance Reviews", level=1)
doc.add_paragraph(
    "Performance reviews are conducted annually in December. Mid-year check-ins are "
    "optional and scheduled at the manager's discretion. Compensation adjustments "
    "take effect at the start of the following fiscal year."
)

doc.save(str(OUT))
print(f"wrote {OUT}")
