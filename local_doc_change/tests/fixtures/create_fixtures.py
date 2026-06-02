"""Deterministic fixture file generator for the local document change pipeline tests.

Run once:  uv run python tests/fixtures/create_fixtures.py

Produces four deterministic files in tests/fixtures/:
- policy.docx   — DOCX with 3 headed sections (requires python-docx)
- handbook.md   — Markdown with 3 H1 sections
- process.txt   — plain text with synthetic SECTION headings
- no_headings.txt — plain text with no headings (fallback-path fixture)
"""

from __future__ import annotations

from pathlib import Path

FIXTURES_DIR = Path(__file__).parent


def _create_policy_docx() -> None:
    """Create a DOCX file with 3 headed sections using python-docx."""
    dest = FIXTURES_DIR / "policy.docx"
    if dest.exists():
        return
    import docx  # noqa: PLC0415

    doc = docx.Document()
    doc.add_heading("Data Retention Policy", level=1)
    doc.add_paragraph(
        "All records must be retained for 7 years per regulatory requirements."
    )
    doc.add_heading("Access Control", level=1)
    doc.add_paragraph(
        "Employee access to production systems requires manager approval."
    )
    doc.add_heading("Incident Response", level=1)
    doc.add_paragraph(
        "Security incidents must be reported within 24 hours of discovery."
    )
    doc.save(str(dest))


def _create_handbook_md() -> None:
    """Create a Markdown file with 3 H1 sections."""
    dest = FIXTURES_DIR / "handbook.md"
    if dest.exists():
        return
    dest.write_text(
        "# Onboarding Process\n\n"
        "New hires must complete orientation within their first week.\n\n"
        "# Benefits Enrollment\n\n"
        "Employees must enroll in benefits within 30 days of start date.\n\n"
        "# Performance Review\n\n"
        "Reviews are conducted annually in December.\n",
        encoding="utf-8",
    )


def _create_process_txt() -> None:
    """Create a plain text file with synthetic SECTION headings."""
    dest = FIXTURES_DIR / "process.txt"
    if dest.exists():
        return
    dest.write_text(
        "SECTION 1. DEPLOYMENT PROCESS\n"
        "Production deployments require two approvals from senior engineers.\n\n"
        "SECTION 2. ROLLBACK PROCEDURE\n"
        "Failed deployments must be rolled back within 15 minutes.\n",
        encoding="utf-8",
    )


def _create_no_headings_txt() -> None:
    """Create a plain text file with no headings for fallback-path testing."""
    dest = FIXTURES_DIR / "no_headings.txt"
    if dest.exists():
        return
    dest.write_text(
        "This document describes general office policies.\n"
        "Employees are expected to maintain professional conduct at all times.\n\n"
        "The company provides health insurance, dental, and vision benefits.\n"
        "Vacation accrual begins after 90 days of employment.\n",
        encoding="utf-8",
    )


def main() -> None:
    """Generate all fixture files idempotently."""
    print(f"Writing fixtures to: {FIXTURES_DIR}")
    _create_policy_docx()
    print("  policy.docx OK")
    _create_handbook_md()
    print("  handbook.md OK")
    _create_process_txt()
    print("  process.txt OK")
    _create_no_headings_txt()
    print("  no_headings.txt OK")
    print("Done.")


if __name__ == "__main__":
    main()
