"""Generate the DOCX files in sample_docs/. Run once; output is committed."""
from pathlib import Path
import docx

HERE = Path(__file__).parent


def build(name, sections):
    d = docx.Document()
    for h, b in sections:
        d.add_heading(h, level=1)
        d.add_paragraph(b)
    d.save(str(HERE / name))
    print("wrote", name)


build("hr_policies.docx", [
    ("Paid Time Off", "Full-time employees accrue 15 days of paid time off per year. PTO must be requested two weeks in advance and approved by the direct manager."),
    ("Remote Work Policy", "Employees are expected to work from the office at least three days per week. Fully remote arrangements require VP approval."),
    ("Expense Reimbursement", "Business expenses under 100 dollars may be approved by a direct manager. Larger expenses require finance approval."),
    ("Performance Reviews", "Performance reviews are conducted annually in December. Compensation adjustments take effect the following fiscal year."),
    ("Parental Leave", "New parents receive 12 weeks of paid parental leave, which may be taken any time within the first year."),
])

build("product_roadmap.docx", [
    ("Supported Platforms", "The Nimbus control app currently supports Windows and macOS. Mobile support is not yet available."),
    ("Launch Schedule", "The Nimbus 2.0 firmware is scheduled to launch on September 15th. The beta program opens one month before launch."),
    ("Hardware Compatibility", "Nimbus 2.0 is compatible with the Atlas and Orion robot arms. Legacy Pioneer arms are supported until end of year."),
])

build("it_policy.docx", [
    ("Device Encryption", "All company laptops must have full-disk encryption enabled before being issued."),
    ("Software Approval", "New software must be approved by IT before installation on company devices."),
    ("Account Provisioning", "New employee accounts are provisioned within two business days of the start date."),
])

print("done")
