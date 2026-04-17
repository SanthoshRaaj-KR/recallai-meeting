from pathlib import Path
from typing import List, Tuple

from bs4 import BeautifulSoup


def document_to_html(path: Path) -> Tuple[str, List[str]]:
    from docx import Document

    document = Document(str(path))
    html_parts: List[str] = []
    headings: List[str] = []
    active_list_kind: str | None = None

    def close_list() -> None:
        nonlocal active_list_kind
        if active_list_kind:
            html_parts.append(f"</{active_list_kind}>")
            active_list_kind = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style_name = (paragraph.style.name if paragraph.style else "").strip()
        if not text:
            close_list()
            continue

        lower_style = style_name.lower()
        if lower_style.startswith("heading"):
            close_list()
            level = 2
            digits = "".join(ch for ch in style_name if ch.isdigit())
            if digits:
                level = max(1, min(int(digits), 6))
            html_parts.append(f"<h{level}>{text}</h{level}>")
            headings.append(text)
            continue

        if "list bullet" in lower_style:
            if active_list_kind != "ul":
                close_list()
                active_list_kind = "ul"
                html_parts.append("<ul>")
            html_parts.append(f"<li>{text}</li>")
            continue

        if "list number" in lower_style:
            if active_list_kind != "ol":
                close_list()
                active_list_kind = "ol"
                html_parts.append("<ol>")
            html_parts.append(f"<li>{text}</li>")
            continue

        close_list()
        html_parts.append(f"<p>{text}</p>")

    close_list()

    for table in document.tables:
        html_parts.append('<table class="officeTable"><tbody>')
        for row_index, row in enumerate(table.rows):
            html_parts.append("<tr>")
            cell_tag = "th" if row_index == 0 else "td"
            for cell in row.cells:
                html_parts.append(f"<{cell_tag}>{cell.text.strip()}</{cell_tag}>")
            html_parts.append("</tr>")
        html_parts.append("</tbody></table>")

    return "".join(html_parts), headings


def html_to_document(html: str, output_path: Path) -> None:
    from docx import Document

    soup = BeautifulSoup(html, "html.parser")
    container = soup.body if soup.body else soup
    document = Document()

    for node in container.children:
        if not getattr(node, "name", None):
            continue
        if node.name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            level = int(node.name[1])
            document.add_heading(node.get_text(" ", strip=True), level=max(0, min(level, 9)))
        elif node.name == "p":
            document.add_paragraph(node.get_text(" ", strip=True))
        elif node.name in {"ul", "ol"}:
            style = "List Bullet" if node.name == "ul" else "List Number"
            for item in node.find_all("li", recursive=False):
                document.add_paragraph(item.get_text(" ", strip=True), style=style)
        elif node.name == "table":
            rows = node.find_all("tr", recursive=False)
            if not rows:
                continue
            first_row = rows[0].find_all(["th", "td"], recursive=False)
            column_count = max(1, len(first_row))
            table = document.add_table(rows=0, cols=column_count)
            for row in rows:
                cells = row.find_all(["th", "td"], recursive=False)
                if not cells:
                    continue
                doc_row = table.add_row().cells
                for index, cell in enumerate(cells[:column_count]):
                    doc_row[index].text = cell.get_text(" ", strip=True)
        else:
            document.add_paragraph(node.get_text(" ", strip=True))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
