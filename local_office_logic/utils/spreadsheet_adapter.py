import csv
from pathlib import Path
from typing import List, Tuple

from bs4 import BeautifulSoup


def workbook_to_html(path: Path) -> Tuple[str, List[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path))
    html_parts: List[str] = []
    targets: List[str] = []
    for worksheet in workbook.worksheets:
        targets.append(worksheet.title)
        html_parts.append(f"<h2>{worksheet.title}</h2>")
        html_parts.append('<table class="officeTable"><tbody>')
        max_row = worksheet.max_row or 1
        max_column = worksheet.max_column or 1
        for row_index in range(1, max_row + 1):
            html_parts.append("<tr>")
            for column_index in range(1, max_column + 1):
                value = worksheet.cell(row=row_index, column=column_index).value
                tag = "th" if row_index == 1 else "td"
                html_parts.append(f"<{tag}>{'' if value is None else value}</{tag}>")
            html_parts.append("</tr>")
        html_parts.append("</tbody></table>")
    return "".join(html_parts), targets


def csv_to_html(path: Path) -> Tuple[str, List[str]]:
    rows: List[List[str]] = []
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle)
        rows = list(reader)

    html_parts: List[str] = [f"<h2>{path.stem}</h2>", '<table class="officeTable"><tbody>']
    for row_index, row in enumerate(rows or [[]]):
        html_parts.append("<tr>")
        tag = "th" if row_index == 0 else "td"
        for cell in row:
            html_parts.append(f"<{tag}>{cell}</{tag}>")
        html_parts.append("</tr>")
    html_parts.append("</tbody></table>")
    return "".join(html_parts), [path.stem]


def html_to_workbook(html: str, output_path: Path) -> None:
    from openpyxl import Workbook

    soup = BeautifulSoup(html, "html.parser")
    container = soup.body if soup.body else soup
    workbook = Workbook()
    first_sheet = workbook.active
    workbook.remove(first_sheet)

    current_title = None
    current_table = None
    for node in container.children:
        if not getattr(node, "name", None):
            continue
        if node.name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            if current_title and current_table is not None:
                _populate_sheet(workbook, current_title, current_table)
            current_title = node.get_text(" ", strip=True) or "Sheet"
            current_table = None
        elif node.name == "table":
            current_table = node

    if current_title and current_table is not None:
        _populate_sheet(workbook, current_title, current_table)
    elif not workbook.worksheets:
        workbook.create_sheet("Sheet1")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(output_path))


def html_to_csv(html: str, output_path: Path) -> None:
    soup = BeautifulSoup(html, "html.parser")
    table = soup.find("table")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        if table:
            for row in table.find_all("tr"):
                writer.writerow([cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])])


def _populate_sheet(workbook, title: str, table_tag) -> None:
    worksheet = workbook.create_sheet(title[:31] or "Sheet")
    for row_index, row in enumerate(table_tag.find_all("tr", recursive=False), start=1):
        for column_index, cell in enumerate(row.find_all(["th", "td"], recursive=False), start=1):
            worksheet.cell(row=row_index, column=column_index, value=cell.get_text(" ", strip=True))
